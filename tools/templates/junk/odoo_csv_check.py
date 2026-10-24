#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Copyright 2026 Zeroincombenze srls <https://www.zeroincombenze.it>
# License LGPL-3.0-or-later (https://www.gnu.org/licenses/lgpl).
"""Controllo di coerenza dei file CSV da importare in Odoo 18.0.

I file han,o nome  {NOME}_[0|1]_{YYYYmmdd}_{HHMM}.csv  con separatore "|";
per ogni {NOME} viene esaminato solo il file più recente.

Le relazioni tra nomi esterni e nomi Odoo vengono da tre file:
- dirmap: {NOME} file -> modello Odoo, colonna chiave (counterpart_pk)
- mapper: colonna del file -> campo Odoo, default, obbligatorietà (req)
- campi Odoo 18.0: modello, campo -> tipo (ttype), modello relazionato

Uso:
    odoo_csv_check.py -c config.ini [-o report.pdf|.md|.xlsx|.ods]
"""
import argparse
import configparser
import csv
import fnmatch
import json
import os
import re
import shutil
import sys
import zipfile
from datetime import date, datetime
from xml.sax.saxutils import escape

try:  # facoltativo: verifica delle partite IVA estere
    from stdnum import get_cc_module
except ImportError:
    get_cc_module = None

__version__ = "18.0.0.1.2"

# Nome file conforme alla specifica (il separatore prima della data può
# essere "_" o "-"); i nomi che non lo rispettano vengono segnalati.
STRICT_PATTERN = re.compile(
    r"^(?P<name>.+)_(?P<flag>[01])[-_](?P<date>\d{8})_(?P<time>\d{4})\.csv$",
    re.IGNORECASE,
)
# Nome file accettato: data e ora facoltative
DEFAULT_PATTERN = (
    r"^(?P<name>.+?)(_[01])?([-_](?P<date>\d{8})(_(?P<time>\d{4}))?)?\.csv$"
)
RE_INTEGER = re.compile(r"^[+-]?\d+$")
RE_FLOAT = re.compile(r"^[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?$")
RE_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
RE_DATETIME = re.compile(
    r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2}(\.\d{1,6})?)?"
    r"(Z|[+-]\d{2}:?\d{2})?$"
)
# Indirizzo e-mail: regola "valid e-mail address" dello standard HTML (WHATWG,
# https://html.spec.whatwg.org/multipage/input.html#valid-e-mail-address),
# con in più l'obbligo di un dominio con TLD (es. .it, .com, .xn--p1ai)
RE_EMAIL = re.compile(
    r"^[a-zA-Z0-9.!#$%&'*+/=?^_`{|}~-]+"
    r"@[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?"
    r"(?:\.[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?)*"
    r"\.(?:[a-zA-Z]{2,63}|xn--[a-zA-Z0-9-]{1,59})$"
)
# Campi Odoo trattati come e-mail: email, *_email, email_* (es. l10n_it_pec_email)
RE_EMAIL_FIELD = re.compile(r"(^|_)email(_|$)")
# Forma "Nome <indirizzo>"
RE_EMAIL_FORMATTED = re.compile(r"^[^<>]*<([^<>]*)>$")
EMAIL_TYPES = ("char", "text")
# Sito web: schema facoltativo (Odoo aggiunge http://), dominio con TLD
# (anche IDN xn--), porta, percorso/query/frammento facoltativi senza spazi;
# impostazione ricavata dalla regex URL di Diego Perini
# (https://gist.github.com/dperini/729294)
RE_WEBSITE = re.compile(
    r"^(?:https?://)?"
    r"(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+"
    r"(?:[a-z]{2,63}|xn--[a-z0-9-]{1,59})"
    r"(?::\d{1,5})?"
    r"(?:[/?#]\S*)?$",
    re.IGNORECASE,
)
# Campi Odoo trattati come sito web: website, *_website, website_*
RE_WEBSITE_FIELD = re.compile(r"(^|_)website(_|$)")
# Ricerca alternativa per e-mail: campo e-mail e nome dei record suggeriti
EMAIL_FIELD = "email"
NAME_FIELD = "name"
MAX_SUGGESTIONS = 10
# Provider di posta pubblici e PEC: lo stesso dominio non indica la stessa
# azienda, quindi la ricerca per dominio non viene fatta (sostituibile con
# public_domains in [ref_search])
PUBLIC_EMAIL_DOMAINS = (
    "alice.it", "blu.it", "email.it", "excite.it", "fastwebnet.it", "iol.it",
    "inwind.it", "jumpy.it", "katamail.com", "libero.it", "poste.it",
    "supereva.it", "tele2.it", "teletu.it", "tim.it", "tin.it", "tiscali.it",
    "tiscalinet.it", "virgilio.it", "vodafone.it",
    "gmail.com", "googlemail.com", "hotmail.com", "hotmail.it", "icloud.com",
    "live.com", "live.it", "mac.com", "me.com", "msn.com", "outlook.com",
    "outlook.it", "yahoo.com", "yahoo.it", "aol.com",
    "bluewin.ch", "free.fr", "gmx.de", "gmx.net", "orange.fr", "t-online.de",
    "wanadoo.fr", "web.de",
    "arubapec.it", "legalmail.it", "pec.it", "postecert.it", "pec.libero.it",
)
# Campi partita IVA e campo del paese usato per validarla (come base_vat)
VAT_FIELDS = ("vat",)
# Codice IVA (posizione fiscale) riportato nell'elenco di [vat_match]
FISCAL_POSITION_FIELD = "property_account_position_id"
COUNTRY_FIELD = "country_id"
RE_VAT_PREFIX = re.compile(r"^([A-Za-z]{2})(.+)$")
# Codice fiscale: l10n_it_codice_fiscale (Odoo), fiscalcode (OCA l10n_it_fiscalcode)
FISCALCODE_FIELDS = ("l10n_it_codice_fiscale", "fiscalcode")
# Persona fisica; con omocodia le cifre possono essere sostituite da LMNPQRSTUV
RE_FISCALCODE = re.compile(
    r"^[A-Z]{6}[0-9LMNPQRSTUV]{2}[ABCDEHLMPRST][0-9LMNPQRSTUV]{2}"
    r"[A-Z][0-9LMNPQRSTUV]{3}[A-Z]$")
FC_OMOCODE = {c: str(ix) for ix, c in enumerate("LMNPQRSTUV")}
FC_MONTHS = "ABCDEHLMPRST"
FC_ODD = dict(zip("0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ",
                  [1, 0, 5, 7, 9, 13, 15, 17, 19, 21] +
                  [1, 0, 5, 7, 9, 13, 15, 17, 19, 21, 2, 4, 18, 20, 11, 3, 6, 8,
                   12, 14, 16, 10, 22, 25, 24, 23]))
# IBAN: lunghezza per paese (registro SWIFT); il BBAN italiano è CIN (lettera),
# ABI (5 cifre), CAB (5 cifre) e conto (12 caratteri alfanumerici)
IBAN_LENGTHS = {
    "AD": 24, "AE": 23, "AL": 28, "AT": 20, "AZ": 28, "BA": 20, "BE": 16,
    "BG": 22, "BH": 22, "BI": 27, "BR": 29, "BY": 28, "CH": 21, "CR": 22,
    "CY": 28, "CZ": 24, "DE": 22, "DJ": 27, "DK": 18, "DO": 28, "EE": 20,
    "EG": 29, "ES": 24, "FI": 18, "FK": 18, "FO": 18, "FR": 27, "GB": 22,
    "GE": 22, "GI": 23, "GL": 18, "GR": 27, "GT": 28, "HN": 28, "HR": 21,
    "HU": 28, "IE": 22, "IL": 23, "IQ": 23, "IS": 26, "IT": 27, "JO": 30,
    "KW": 30, "KZ": 20, "LB": 28, "LC": 32, "LI": 21, "LT": 20, "LU": 20,
    "LV": 21, "LY": 25, "MC": 27, "MD": 24, "ME": 22, "MK": 19, "MN": 20,
    "MR": 27, "MT": 31, "MU": 30, "NI": 28, "NL": 18, "NO": 15, "OM": 23,
    "PK": 24, "PL": 28, "PS": 29, "PT": 25, "QA": 29, "RO": 24, "RS": 22,
    "RU": 33, "SA": 24, "SC": 31, "SD": 18, "SE": 24, "SI": 19, "SK": 24,
    "SM": 27, "SO": 23, "ST": 25, "SV": 28, "TL": 23, "TN": 24, "TR": 26,
    "UA": 29, "VA": 22, "VG": 24, "XK": 20, "YE": 30,
}
RE_IBAN = re.compile(r"^[A-Z]{2}\d{2}[A-Z0-9]+$")
RE_IBAN_BBAN = {
    "IT": re.compile(r"^[A-Z]\d{10}[A-Z0-9]{12}$"),
    "SM": re.compile(r"^[A-Z]\d{10}[A-Z0-9]{12}$"),
}
RE_QUOTED = re.compile(r'^"((?:[^"]|"")*)"$', re.DOTALL)
RE_REF_TOKEN = re.compile(r'\s*(?:"((?:[^"]|"")*)"|\'([^\']*)\'|([^,]*))\s*(?:,|$)')

INTEGER_TYPES = ("integer",)
FLOAT_TYPES = ("float", "monetary")
REF_TYPES = ("many2one", "one2many", "many2many")
X2MANY_TYPES = ("one2many", "many2many")
MAX_RECORD_LINES = 50
TRUE_VALUES = ("1", "true", "t", "yes", "y", "x")

# Intestazioni accettate nel file dei campi Odoo (minuscolo) -> chiave interna;
# "model", "name", "ttype", "relation" sono quelli di ir.model.fields
FIELD_HEADERS = {
    "model": "model",
    "modello": "model",
    "model name": "model",
    "name": "field",
    "field": "field",
    "field name": "field",
    "campo": "field",
    "nome campo": "field",
    "ttype": "ttype",
    "type": "ttype",
    "field type": "ttype",
    "tipo": "ttype",
    "tipo campo": "ttype",
    "relation": "relation",
    "object relation": "relation",
    "modello relazionato": "relation",
    "relazione": "relation",
    "selection": "selection",
    "valori": "selection",
}


class Issue:
    """Segnalazione; detail è un'informazione aggiuntiva (es. i record
    suggeriti) che nel PDF non viene raggruppata né abbreviata"""
    __slots__ = ("row", "field", "base", "detail", "message", "value", "searched",
                 "ref", "pinned")

    def __init__(self, row, field, message, detail="", value="", searched=False):
        self.row = row
        self.field = field
        self.base = message
        self.detail = detail
        self.value = value  # valore della cella segnalata
        self.searched = searched  # detail viene dalla ricerca alternativa
        self.message = "%s; %s" % (message, detail) if detail else message
        self.ref = ""  # codice del record ([row_ref]) per correggere l'origine
        self.pinned = False  # colonna di [row_ref]: codice anche nei gruppi del PDF

    @property
    def label(self):
        """Riga con il codice del record, se noto: 198 (C  1904)"""
        return "%s (%s)" % (self.row, self.ref) if self.ref else str(self.row)


class FileResult:
    def __init__(self, name, path, model):
        self.name = name
        self.path = path
        self.model = model
        self.records = 0
        self.bad_records = 0
        self.issues = []
        self.warned = set()  # segnalazioni da emettere una sola volta

    @property
    def valid_pct(self):
        return pct(self.records, self.bad_records)


def pct(total, bad):
    if not total:
        return None
    return round((total - bad) * 100.0 / total, 2)


def fmt_pct(value):
    return "n/d" if value is None else "%.2f%%" % value


def split_list(value):
    return [x.strip() for x in re.split(r"[,;\s]+", value or "") if x.strip()]


def unquote(value):
    """Toglie i doppi apici attorno all'intero valore ("" -> ")"""
    value = value.strip()
    match = RE_QUOTED.match(value)
    if match:
        value = match.group(1).replace('""', '"')
    return value


def split_refs(value):
    """Separa i riferimenti multipli (x2many) su virgola; ogni riferimento
    può essere racchiuso tra doppi apici o apici singoli"""
    refs = []
    pos = 0
    while pos < len(value):
        match = RE_REF_TOKEN.match(value, pos)
        if not match or match.end() == pos:
            break
        dq, sq, bare = match.groups()
        ref = dq.replace('""', '"') if dq is not None else sq if sq is not None else bare
        if ref and ref.strip():
            refs.append(ref.strip())
        pos = match.end()
    return refs


def normalize_header(name):
    """'partner_id/id' o 'partner_id/.id' -> 'partner_id'"""
    name = unquote(name)
    for suffix in ("/.id", "/id", ":id"):
        if name.endswith(suffix):
            return name[: -len(suffix)]
    return name


# ---------------------------------------------------------------------------
# Configurazione
# ---------------------------------------------------------------------------
class Config:
    def __init__(self, path):
        if not os.path.isfile(path):
            raise SystemExit("File di configurazione %s non trovato" % path)
        cp = configparser.ConfigParser(interpolation=None)
        cp.optionxform = str  # i nomi file/modello sono case sensitive
        cp.read(path, encoding="utf-8")
        base = os.path.dirname(os.path.abspath(path))

        def get_path(key, default=None):
            value = cp.get("main", key, fallback=default)
            if value:
                value = os.path.expanduser(value)
                if not os.path.isabs(value):
                    value = os.path.normpath(os.path.join(base, value))
            return value

        if not cp.has_section("main"):
            raise SystemExit("Sezione [main] mancante in %s" % path)
        self.data_dir = get_path("data_dir")
        self.dirmap_file = get_path("dirmap_file", "odoo_csv_check_dirmap.csv")
        self.mapper_file = get_path("mapper_file", "odoo_csv_check_mapper.csv")
        self.fields_file = get_path("fields_file", "odoo_csv_check_fields.csv")
        self.fields_sheet = cp.get("main", "fields_sheet", fallback="") or None
        self.backend = cp.get("main", "backend", fallback="").strip()
        self.report_file = get_path("report_file", "report_%(datetime)s.md")
        self.encoding = cp.get("main", "encoding", fallback="utf-8-sig")
        self.file_pattern = re.compile(
            cp.get("main", "file_pattern", fallback="") or DEFAULT_PATTERN,
            re.IGNORECASE)
        self.ref_ignore_case = cp.getboolean(
            "main", "ref_ignore_case", fallback=True)
        self.prefix_match = cp.getboolean("main", "prefix_match", fallback=True)
        if not self.data_dir or not os.path.isdir(self.data_dir):
            raise SystemExit("Directory dati '%s' non trovata" % self.data_dir)
        for label, fn in (("dirmap", self.dirmap_file),
                          ("mapper", self.mapper_file),
                          ("campi Odoo", self.fields_file)):
            if not fn or not os.path.isfile(fn):
                raise SystemExit("File %s '%s' non trovato" % (label, fn))
        # [keys] NOME o modello = colonne chiave aggiuntive a counterpart_pk
        keys = dict(cp.items("keys")) if cp.has_section("keys") else {}
        self.keys = {k: split_list(v) for k, v in keys.items()}
        # [external] models = modelli già presenti in Odoo: riferimenti non verificati
        self.external_models = set(
            split_list(cp.get("external", "models", fallback=""))
        )
        # [equivalent] modello = altri modelli le cui chiavi valgono anche per esso
        self.equivalent = {}
        if cp.has_section("equivalent"):
            for model, others in cp.items("equivalent"):
                self.equivalent[model] = split_list(others)
        # [hierarchy] modello = separatore: le chiavi sono percorsi (es. il
        # complete_name "A / B / C" di product.category) e anche ogni loro
        # prefisso ("A", "A / B") è un riferimento valido
        self.hierarchy = {}
        if cp.has_section("hierarchy"):
            for model, sep in cp.items("hierarchy"):
                sep = sep.strip().strip('"')
                self.hierarchy[model] = sep.center(len(sep) + 2) if sep else " / "
        # [selection] modello.campo = valori validi delle selection dinamiche
        # (calcolate a runtime, senza valori nel file dei campi)
        # [ignore] {NOME} = colonne da non controllare (restano note, quindi
        # non vengono segnalate come assenti dal mapper)
        self.ignore = {}
        if cp.has_section("ignore"):
            for name, cols in cp.items("ignore"):
                self.ignore[name] = split_list(cols)
        # [required] modello = campi che devono avere un valore in ogni riga
        # dei file di quel modello (oltre ai req del mapper)
        self.required = {}
        if cp.has_section("required"):
            for model, fields in cp.items("required"):
                self.required[model] = split_list(fields)
        # [ref_files] modello.campo = {NOME} dei soli file le cui chiavi sono
        # valide per quel riferimento (gli altri file del modello, es. le
        # estensioni dell'anagrafica, non contano); {NOME}.campo vale solo per
        # il file con quel nome esatto e prevale su modello.campo
        self.ref_files = {}
        if cp.has_section("ref_files"):
            for key, names in cp.items("ref_files"):
                model, _sep, field = key.rpartition(".")
                if model and field:
                    self.ref_files[(model, field)] = split_list(names)
        # [ref_suffix] modello collegato = testo con cui devono terminare i
        # riferimenti a quel modello (es. product.category = .); exclude =
        # {NOME} dei file esclusi dalla regola
        self.ref_suffix = {}
        self.ref_suffix_exclude = set()
        if cp.has_section("ref_suffix"):
            for key, value in cp.items("ref_suffix"):
                if key == "exclude":
                    self.ref_suffix_exclude = set(split_list(value))
                else:
                    self.ref_suffix[key] = value.strip().strip('"')
        # [ref_insert] modello.campo o {NOME}.campo = {NOME} del file del
        # modello collegato: i riferimenti mancanti non sono errori del file
        # ma record da inserire in quel file (es. i tag res.partner.category)
        self.ref_insert = {}
        if cp.has_section("ref_insert"):
            for key, target in cp.items("ref_insert"):
                model, _sep, field = key.rpartition(".")
                if model and field:
                    self.ref_insert[(model, field)] = target.strip()
        # [ref_child_email] {NOME}.campo o modello.campo = campo che lega il
        # figlio al padre (es. parent_id): i valori sono e-mail di contatti
        # figli del partner della riga, cercati per (padre, e-mail) come fa
        # l'importatore (es. ACA -> extra_sale_partner_ids)
        self.ref_child_email = {}
        if cp.has_section("ref_child_email"):
            for key, parent_field in cp.items("ref_child_email"):
                model, _sep, field = key.rpartition(".")
                if model and field:
                    self.ref_child_email[(model, field)] = parent_field.strip() or "parent_id"
        # [ref_search] modello.campo = {NOME} dei file di anagrafica: se il
        # riferimento non è trovato, cerca in quei file l'e-mail e la partita
        # IVA della riga ([ref_email] è il vecchio nome della sezione)
        self.ref_search = {}
        self.public_domains = set(PUBLIC_EMAIL_DOMAINS)
        section = next((x for x in ("ref_search", "ref_email") if cp.has_section(x)), None)
        if section:
            for key, names in cp.items(section):
                if key == "public_domains":
                    # domini di provider pubblici: niente ricerca per dominio
                    self.public_domains = {d.lower() for d in split_list(names)}
                    continue
                model, _sep, field = key.rpartition(".")
                if model and field:
                    self.ref_search[(model, field)] = split_list(names)
        # [ref_values] modello.campo o {NOME}.campo = valori sempre validi per
        # il riferimento anche se assenti dai file dati (es. listini 1, 2, 3
        # già presenti in Odoo)
        self.ref_values = {}
        if cp.has_section("ref_values"):
            for key, values in cp.items("ref_values"):
                model, _sep, field = key.rpartition(".")
                if model and field:
                    self.ref_values[(model, field)] = split_list(values)
        # [unique_only] {NOME} = colonne: il file viene controllato solo per
        # i valori ripetuti in ciascuna colonna (nessun altro controllo)
        self.unique_only = {}
        if cp.has_section("unique_only"):
            for name, cols in cp.items("unique_only"):
                self.unique_only[name] = split_list(cols)
        # [lookup_only] {NOME}.colonna = file: colonna, oppure = modello: il
        # file viene controllato solo per i valori della colonna, che devono
        # comparire nella colonna dell'altro file (nome file o {NOME} in
        # data_dir, separatore | ; , o tab) o tra le chiavi del modello
        self.lookup_only = {}
        if cp.has_section("lookup_only"):
            for key, value in cp.items("lookup_only"):
                name, _sep, col = key.rpartition(".")
                source, sep, target = value.partition(":")
                if name and col and source.strip():
                    self.lookup_only.setdefault(name, []).append(
                        (col, source.strip(), target.strip() if sep else None))
        # [alias] pattern = {NOME} di dirmap: i file il cui {NOME} corrisponde
        # al pattern (glob, come filename_pattern dell'importatore) usano quella
        # mappatura; vale prima di prefix_match
        self.alias = []
        if cp.has_section("alias"):
            for pattern, name in cp.items("alias"):
                self.alias.append((pattern.upper(), name.strip()))
        # [row_ref] {NOME} = colonna codice: colonne; gli errori sulle colonne
        # indicate riportano accanto alla riga il codice del record
        self.row_ref = {}
        if cp.has_section("row_ref"):
            for name, value in cp.items("row_ref"):
                code, _sep, cols = value.partition(":")
                if code.strip() and cols.strip():
                    self.row_ref[name] = (code.strip(), split_list(cols))
        # [unique] {NOME} = colonne: nessuna riga con la stessa combinazione
        # di valori di tutte le colonne indicate (oltre agli altri controlli)
        self.unique = {}
        if cp.has_section("unique"):
            for name, cols in cp.items("unique"):
                self.unique[name] = split_list(cols)
        # [requires] {NOME}.colonna = colonne che devono avere un valore nelle
        # righe in cui la colonna ha un valore
        self.requires = {}
        if cp.has_section("requires"):
            for key, cols in cp.items("requires"):
                name, _sep, col = key.rpartition(".")
                if name and col:
                    self.requires.setdefault(name, {})[col] = split_list(cols)
        # [vat_match] {NOME} = altri {NOME}: elenco dei record con la stessa
        # partita IVA nei due file (es. cliente e fornitore importati come due
        # partner distinti); non sono errori, hanno una sezione propria
        self.vat_match = {}
        if cp.has_section("vat_match"):
            for name, others in cp.items("vat_match"):
                self.vat_match[name] = split_list(others)
        # [purge] {NOME} = data (YYYY-MM-DD): funzione temporanea, solo in
        # quella data i record con errori vengono tolti dal file (l'originale
        # resta in .bak)
        self.purge = {}
        if cp.has_section("purge"):
            for name, day in cp.items("purge"):
                self.purge[name] = day.strip()
        # [iban] {NOME} = colonna IBAN: prima dei controlli il file viene
        # riscritto (l'originale resta in .bak) ricalcolando le cifre di
        # controllo "00" degli IBAN fittizi (es. IT00A...) e togliendo le
        # righe con IBAN non valido
        self.iban = {}
        if cp.has_section("iban"):
            for name, col in cp.items("iban"):
                if col.strip():
                    self.iban[name] = normalize_header(col.strip())
        self.selections = {}
        if cp.has_section("selection"):
            for key, values in cp.items("selection"):
                model, _sep, field = key.rpartition(".")
                if model and field:
                    self.selections[(model, field)] = {
                        v: "" for v in split_list(values)}


# ---------------------------------------------------------------------------
# Mappature (dirmap, mapper) e campi Odoo
# ---------------------------------------------------------------------------
def read_table(path, sheet=None, encoding="utf-8-sig"):
    """Restituisce le righe (liste di stringhe) di un file csv o xlsx"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            raise SystemExit(
                "Per leggere %s serve openpyxl (pip install openpyxl)" % path
            )
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return [
            ["" if c is None else str(c).strip() for c in row]
            for row in ws.iter_rows(values_only=True)
        ]
    if ext == ".csv":
        with open(path, encoding=encoding, newline="") as fd:
            sample = fd.read(4096)
            fd.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;|\t")
            except csv.Error:
                dialect = csv.excel
            return [[c.strip() for c in row] for row in csv.reader(fd, dialect)]
    raise SystemExit("Formato file non supportato: %s (usare csv o xlsx)" % path)


def read_dicts(path, encoding="utf-8-sig"):
    """Righe di un csv come dizionari con chiavi e valori ripuliti dagli spazi"""
    rows = read_table(path, encoding=encoding)
    if not rows:
        return []
    header = [h.strip() for h in rows[0]]
    return [
        {k: (row[ix].strip() if ix < len(row) else "") for ix, k in enumerate(header)}
        for row in rows[1:] if any(row)
    ]


class Dirmap:
    """Un {NOME} file: modello Odoo e colonna chiave"""
    __slots__ = ("name", "model", "pk", "note")

    def __init__(self, name, model, pk, note):
        self.name = name
        self.model = model
        self.pk = pk
        self.note = note


class Mapping:
    """Tipo di un campo mappato: una colonna del file {NOME}"""
    __slots__ = ("column", "field", "ttype", "relation", "default", "apply4",
                 "required", "strict", "selection")

    def __init__(self, column, field, ttype, relation, default, apply4, required,
                 selection=None):
        self.column = column
        self.field = field
        self.ttype = ttype
        self.relation = relation
        self.default = default
        self.apply4 = apply4
        self.required = required  # motivo dell'obbligo ("" = facoltativo)
        self.strict = False  # True: il valore deve essere nel file (no default)
        self.selection = selection  # {chiave: etichetta} o None se non noti

    @property
    def needs_value(self):
        """Il file deve fornire un valore: obbligatorio e senza default che
        lo sostituisca (per [required] il default non basta)"""
        return bool(self.required) and (self.strict or not self.default)


class Mappings:
    """Carica odoo_csv_check_dirmap.csv, odoo_csv_check_mapper.csv e le
    definizioni dei campi Odoo (odoo_csv_check_fields.csv).

    - dirmaps: {NOME: Dirmap}
    - columns: {NOME: {colonna: Mapping}}; Mapping.field vuoto = colonna
      nota ma non importata
    - problems: {NOME: [messaggi]} per incongruenze tra mapper e Odoo
    """

    def __init__(self, cfg):
        self.fields = self.load_fields(cfg)
        self.dirmaps = {}
        for row in read_dicts(cfg.dirmap_file):
            if row.get("type", "dirmap") != "dirmap":
                continue
            if cfg.backend and row.get("backend") != cfg.backend:
                continue
            name, model = row.get("ext_name"), row.get("loc_name")
            if name and model:
                self.dirmaps[name] = Dirmap(
                    name, model, row.get("counterpart_pk", ""), row.get("note", ""))
        self.columns = {}
        self.problems = {}
        for row in read_dicts(cfg.mapper_file):
            if row.get("type", "mapper") != "mapper":
                continue
            if cfg.backend and row.get("backend") != cfg.backend:
                continue
            name, column = row.get("dirmap"), row.get("ext_name")
            field = row.get("loc_name", "")
            if not name or not column:
                continue  # valore solo lato Odoo (default): nulla da controllare
            dirmap = self.dirmaps.get(name)
            if not dirmap:
                self.problems.setdefault(name, []).append(
                    "mapper: {NOME} '%s' assente in dirmap" % name)
                continue
            ttype = relation = ""
            selection = None
            if field:
                info = self.fields.get(dirmap.model, {}).get(field)
                if info:
                    ttype, relation, selection = info
                    selection = cfg.selections.get((dirmap.model, field), selection)
                else:
                    self.problems.setdefault(name, []).append(
                        "mapper: campo '%s' (colonna %s) inesistente nel modello %s"
                        % (field, column, dirmap.model))
            cols = self.columns.setdefault(name, {})
            if column in cols and cols[column].field:
                continue  # prima definizione valida
            cols[column] = Mapping(
                column, field, ttype, relation, row.get("default", ""),
                row.get("apply4", ""),
                "req nel mapper" if row.get("req", "").lower() in TRUE_VALUES else "",
                selection)
        # colonne escluse dal controllo ([ignore]): note ma non importate
        for name, cols in cfg.ignore.items():
            for col in cols:
                self.columns.setdefault(name, {})[col] = Mapping(
                    col, "", "", "", "", "", "")
        # campi obbligatori per modello da [required] della configurazione
        for name, dirmap in self.dirmaps.items():
            for field in cfg.required.get(dirmap.model, []):
                cols = [mp for mp in self.columns.get(name, {}).values()
                        if mp.field == field]
                if not cols:
                    self.problems.setdefault(name, []).append(
                        "mapper: nessuna colonna per il campo obbligatorio %s "
                        "([required] %s)" % (field, dirmap.model))
                for mapping in cols:
                    # prevale su req del mapper: il default non basta
                    mapping.required = "[required] %s" % dirmap.model
                    mapping.strict = True

    @staticmethod
    def load_fields(cfg):
        """{modello: {campo: (ttype, relation, selection)}}; selection è il
        dizionario {chiave: etichetta} dei valori validi, None se non noti"""
        rows = [r for r in read_table(cfg.fields_file, cfg.fields_sheet, cfg.encoding)
                if any(r)]
        if not rows:
            raise SystemExit("File campi Odoo %s vuoto" % cfg.fields_file)
        colmap = {}
        for ix, head in enumerate(rows[0]):
            key = FIELD_HEADERS.get(head.strip().lower())
            if key and key not in colmap:
                colmap[key] = ix
        missing = {"model", "field", "ttype"} - set(colmap)
        if missing:
            raise SystemExit(
                "File campi Odoo %s: colonne mancanti %s"
                % (cfg.fields_file, ", ".join(sorted(missing))))

        def cell(row, key):
            ix = colmap.get(key)
            return row[ix].strip() if ix is not None and ix < len(row) else ""

        fields = {}
        for row in rows[1:]:
            model, field = cell(row, "model"), cell(row, "field")
            if model and field:
                fields.setdefault(model, {})[field] = (
                    cell(row, "ttype").lower(), cell(row, "relation"),
                    parse_selection(cell(row, "selection")))
        return fields

    def resolve(self, name, prefix_match=True, alias=()):
        """Restituisce (chiave dirmap, associato per alias o prefisso) per un
        {NOME}"""
        if name in self.dirmaps:
            return name, False
        for pattern, key in alias:
            if key in self.dirmaps and fnmatch.fnmatchcase(name.upper(), pattern):
                return key, True
        if prefix_match:
            candidates = [d for d in self.dirmaps if name.startswith(d + "_")]
            if candidates:
                return max(candidates, key=len), True
        return None, False


# ---------------------------------------------------------------------------
# Selezione e lettura file dati
# ---------------------------------------------------------------------------
def select_files(cfg):
    """Restituisce ({NOME: path del più recente}, [file scartati], {NOME})
    l'ultimo elemento contiene i NOME il cui file non rispetta la specifica"""
    latest = {}
    skipped = []
    for fn in sorted(os.listdir(cfg.data_dir)):
        path = os.path.join(cfg.data_dir, fn)
        if not os.path.isfile(path):
            continue
        match = cfg.file_pattern.match(fn)
        if not match:
            if fn.lower().endswith(".csv"):
                skipped.append((fn, "nome non conforme al pattern"))
            continue
        groups = match.groupdict()
        stamp = (groups.get("date") or "") + (groups.get("time") or "")
        name = groups["name"]
        if name in latest:
            old_stamp, old_fn = latest[name]
            if stamp > old_stamp:
                skipped.append((old_fn, "superato da %s" % fn))
            else:
                skipped.append((fn, "superato da %s" % old_fn))
                continue
        latest[name] = (stamp, fn)
    nonconforming = {
        name for name, (_s, fn) in latest.items() if not STRICT_PATTERN.match(fn)}
    return (
        {name: os.path.join(cfg.data_dir, fn) for name, (_s, fn) in latest.items()},
        skipped,
        nonconforming,
    )


def split_record(text):
    """Divide un record su '|' fuori dai doppi apici"""
    fields = []
    buf = []
    in_quote = False
    for ch in text:
        if ch == '"':
            in_quote = not in_quote
        elif ch == "|" and not in_quote:
            fields.append("".join(buf))
            buf = []
            continue
        buf.append(ch)
    fields.append("".join(buf))
    return fields


def read_data_file(path, encoding):
    """Genera (numero riga, lista valori grezzi, errore); il primo record è
    l'intestazione.

    Il numero riga è quello del foglio di calcolo (LibreOffice, Excel):
    l'intestazione è la riga 1, ogni record è una riga anche se va a capo
    dentro un campo e anche le righe vuote contano.

    Un campo racchiuso tra doppi apici può contenere '|' e andare a capo:
    il record prosegue sulle righe successive finché gli apici non sono
    bilanciati. Oltre MAX_RECORD_LINES righe il record viene considerato
    errato (apici non chiusi) e la lettura riprende dalla riga seguente.
    """
    with open(path, encoding=encoding, newline="") as fd:
        text = fd.read()
    # solo \r\n, \n e \r chiudono una riga (splitlines spezza anche su
    # \x0b, \x0c, \x85, \u2028... che un foglio di calcolo non considera)
    lines = re.split(r"\r\n|\n|\r", text)
    if lines and lines[-1] == "":
        lines.pop()
    ix = 0
    rownum = 0
    while ix < len(lines):
        start = ix
        rownum += 1
        if not lines[ix].strip():
            ix += 1
            continue
        text = lines[ix]
        quotes = text.count('"')
        ix += 1
        while quotes % 2 and ix < len(lines) and ix - start < MAX_RECORD_LINES:
            text += "\n" + lines[ix]
            quotes += lines[ix].count('"')
            ix += 1
        if quotes % 2:
            ix = start + 1
            yield rownum, lines[start].split("|"), "Doppi apici non bilanciati"
            continue
        yield rownum, split_record(text), None


def match_header(header, column):
    """Indice della colonna nell'intestazione, senza distinzione
    maiuscole/minuscole e spazi ripetuti; None se assente"""
    wanted = " ".join(column.split()).lower()
    for ix, col in enumerate(header):
        if " ".join(col.split()).lower() == wanted:
            return ix
    return None


def read_column(path, column, encoding):
    """Valori non vuoti di una colonna di un file csv con separatore | ; ,
    o tab (il più frequente nella prima riga); None se la colonna manca"""
    with open(path, encoding=encoding, newline="") as fd:
        first = fd.readline()
        fd.seek(0)
        delimiter = max("|;,\t", key=first.count)
        rows = csv.reader(fd, delimiter=delimiter)
        header = [normalize_header(h) for h in next(rows, [])]
        ix = match_header(header, column)
        if ix is None:
            return None
        return [row[ix].strip() for row in rows if ix < len(row) and row[ix].strip()]


def lookup_sources(cfg, files):
    """{NOME} dei file usati come elenco di valori da [lookup_only]"""
    sources = {src for rules in cfg.lookup_only.values()
               for _col, src, target in rules if target is not None}
    return {name for name, path in files.items()
            if name in sources or os.path.basename(path) in sources}


# ---------------------------------------------------------------------------
# Controlli
# ---------------------------------------------------------------------------
class Checker:
    def __init__(self, cfg, maps, files, nonconforming=()):
        self.cfg = cfg
        self.maps = maps
        self.files = files
        self.nonconforming = set(nonconforming)
        self.ref_keys = {}  # modello -> set di chiavi valide
        self.ref_files = {}  # modello -> {NOME} dei file che forniscono le chiavi
        self.name_keys = {}  # {NOME} (dirmap) -> set di chiavi del file
        self._only_keys = {}  # cache: tupla di {NOME} -> unione delle chiavi
        # [ref_insert]: {NOME destinazione: {valore: {NOME origine: [righe]}}}
        self.to_insert = {}
        self._insert_model = {}  # NOME destinazione -> modello
        # [ref_child_email]: campo padre -> {(chiave padre, e-mail)}, e i file letti
        self.child_emails = {}
        self.child_files = {}
        # [ref_search]: {NOME: {"email"|"domain"|"vat": {valore: [record]}}}
        self.alt_index = {}
        self._row = None  # riga in esame: {colonna: valore} e mappature
        self._models_with_data = set()
        self.lookup_files = {}  # [lookup_only]: {NOME} -> file di soli valori
        self.iban_issues = {}  # [iban]: {NOME} -> segnalazioni della correzione
        self.dirmap_of = {
            name: maps.resolve(name, cfg.prefix_match, cfg.alias) for name in files}

    def model_of(self, name):
        key = self.dirmap_of[name][0]
        return self.maps.dirmaps[key].model if key else None

    def key_columns(self, name, model):
        key = self.dirmap_of[name][0]
        cols = split_list(self.maps.dirmaps[key].pk) if key else []
        for extra in (self.cfg.keys.get(model, []), self.cfg.keys.get(key, []),
                      self.cfg.keys.get(name, [])):
            cols += [c for c in extra if c not in cols]
        return cols

    def option(self, section, result, field):
        """Valore di una sezione per {NOME}.campo: nome del file, poi {NOME}
        di dirmap a cui è associato (alias o prefisso), poi modello.campo"""
        key = self.dirmap_of[result.name][0]
        for owner in (result.name, key, result.model):
            if owner and (owner, field) in section:
                return section[(owner, field)]
        return None

    def norm_key(self, value):
        value = value.strip()
        return value.upper() if self.cfg.ref_ignore_case else value

    def collect_keys(self):
        """Carica, per ogni modello, i valori delle colonne chiave di tutti i
        suoi file: sono il dominio valido dei riferimenti."""
        for name, path in self.files.items():
            model = self.model_of(name)
            if not model:
                continue
            self._models_with_data.add(model)
            files = self.ref_files.setdefault(model, [])
            files.append(name)  # nome breve: {NOME} senza flag, data e ora
            keys = self.ref_keys.setdefault(model, set())
            own = self.name_keys.setdefault(self.dirmap_of[name][0], set())
            key_cols = self.key_columns(name, model)
            sep = self.cfg.hierarchy.get(model)
            header = None
            for _lineno, values, error in read_data_file(path, self.cfg.encoding):
                if header is None:
                    header = [normalize_header(h) for h in values]
                    ixs = [header.index(k) for k in key_cols if k in header]
                    continue
                if error:
                    continue
                for ix in ixs:
                    if ix < len(values):
                        value = unquote(values[ix])
                        if value:
                            found = {self.norm_key(value)}
                            if sep:
                                parts = value.split(sep)
                                for n in range(1, len(parts)):
                                    found.add(self.norm_key(sep.join(parts[:n])))
                            keys |= found
                            own |= found
        self.collect_alternatives()
        self.collect_child_emails()
        # modelli equivalenti (es. product.product <-> product.template)
        for model, others in self.cfg.equivalent.items():
            for other in others:
                if other in self.ref_keys:
                    self.ref_keys.setdefault(model, set()).update(self.ref_keys[other])
                    self._models_with_data.add(model)
                    files = self.ref_files.setdefault(model, [])
                    files += [f for f in self.ref_files.get(other, [])
                              if f not in files]

    def collect_alternatives(self):
        """Indici dei file di anagrafica indicati in [ref_search] per la
        ricerca alternativa: e-mail esatta, dominio e partita IVA.
        Ogni record è (NOME, riga, chiave, nome)"""
        wanted = {n for names in self.cfg.ref_search.values() for n in names}
        for name, path in self.files.items():
            key = self.dirmap_of[name][0]
            if key not in wanted:
                continue
            cols = self.maps.columns.get(key, {})
            email_cols = [c for c, mp in cols.items() if mp.field == EMAIL_FIELD]
            name_cols = [c for c, mp in cols.items() if mp.field == NAME_FIELD]
            vat_cols = [c for c, mp in cols.items()
                        if mp.field in VAT_FIELDS + FISCALCODE_FIELDS]
            key_cols = split_list(self.maps.dirmaps[key].pk)
            index = self.alt_index.setdefault(
                key, {"email": {}, "domain": {}, "vat": {}})
            header = None
            for rownum, values, error in read_data_file(path, self.cfg.encoding):
                if header is None:
                    header = [normalize_header(h) for h in values]
                    continue
                if error or len(values) != len(header):
                    continue
                row = dict(zip(header, (unquote(v) for v in values)))
                record = (key, rownum,
                          " / ".join(row.get(c, "") for c in key_cols if row.get(c)),
                          " ".join(row.get(c, "") for c in name_cols if row.get(c)))
                for col in email_cols:
                    for address in email_addresses(row.get(col, "")):
                        index["email"].setdefault(address, []).append(record)
                        index["domain"].setdefault(
                            address.split("@")[1], []).append(record)
                for vat in {normalize_vat(row.get(c, "")) for c in vat_cols} - {""}:
                    index["vat"].setdefault(vat, []).append(record)

    def read_vat_rows(self, wanted):
        """Righe del file {NOME} (nome o NOME di dirmap) con partita IVA:
        (colonna chiave, colonna IVA, colonna nome, colonna codice IVA,
        [(chiave, IVA, IVA normalizzata, nome, codice IVA)]); None se il file
        non c'è"""
        name = next((n for n in sorted(self.files) if n == wanted), None) or next(
            (n for n in sorted(self.files) if self.dirmap_of[n][0] == wanted), None)
        if not name:
            return None
        key = self.dirmap_of[name][0]
        cols = self.maps.columns.get(key, {}) if key else {}
        pk_cols = split_list(self.maps.dirmaps[key].pk) if key else []
        id_col = pk_cols[0] if pk_cols else ""
        vat_col = next((c for c, mp in cols.items() if mp.field in VAT_FIELDS), "")
        name_col = next((c for c, mp in cols.items() if mp.field == NAME_FIELD), "")
        fpos_col = next(
            (c for c, mp in cols.items() if mp.field == FISCAL_POSITION_FIELD), "")
        rows = []
        header = None
        for _lineno, values, error in read_data_file(self.files[name], self.cfg.encoding):
            if header is None:
                header = [normalize_header(h) for h in values]
                continue
            if error or len(values) != len(header):
                continue
            row = dict(zip(header, (unquote(v) for v in values)))
            vat = normalize_vat(row.get(vat_col, ""))
            if vat:
                rows.append((row.get(id_col, ""), row[vat_col], vat,
                             row.get(name_col, ""), row.get(fpos_col, "")))
        return id_col, vat_col, name_col, fpos_col, rows

    def vat_matches(self):
        """[vat_match]: tabelle [(titolo, intestazione, righe)] dei record
        con la stessa partita IVA: chiave, IVA e codice IVA dei due file,
        nome del primo"""
        tables = []
        for name, others in self.cfg.vat_match.items():
            source = self.read_vat_rows(name)
            if source is None:
                continue
            id_col, vat_col, name_col, fpos_col, rows = source
            for other in others:
                target = self.read_vat_rows(other)
                if target is None:
                    continue
                index = {}
                for rec in target[4]:
                    index.setdefault(rec[2], []).append(rec)
                header = ["%s %s" % (name, id_col), "%s %s" % (name, vat_col),
                          "%s %s" % (name, fpos_col),
                          "%s %s" % (other, target[0]), "%s %s" % (other, target[1]),
                          "%s %s" % (other, target[3]),
                          "%s %s" % (name, name_col)]
                found = [[code, vat, fpos, ocode, ovat, ofpos, desc]
                         for code, vat, norm, desc, fpos in rows
                         for ocode, ovat, _norm, _desc, ofpos in index.get(norm, [])]
                tables.append(("Partita IVA in %s e %s" % (name, other), header, found))
        return tables

    def add_insert_issues(self, results):
        """Aggiunge al file collegato ([ref_insert]) i record da inserire, uno
        per valore, con i file e le righe che lo usano; se il file non c'è
        viene aggiunta una sezione apposita"""
        by_key = {self.dirmap_of[r.name][0] or r.name: r for r in results}
        for target, values in sorted(self.to_insert.items()):
            model = self._insert_model.get(target, "")
            result = by_key.get(target)
            if result is None:
                result = FileResult(target, "%s (file assente)" % target, model)
                results.append(result)
            dirmap = self.maps.dirmaps.get(target)
            field = dirmap.pk if dirmap and dirmap.pk else "-"
            for value, sources in sorted(values.items()):
                uses = []
                for name, rows in sorted(sources.items()):
                    rows_text = compress_rows(rows)
                    parts = rows_text.split(", ")
                    if len(parts) > 12:
                        rows_text = ", ".join(parts[:12]) + ", …"
                    uses.append("%s: %d righe (%s)" % (name, len(rows), rows_text))
                result.issues.append(Issue(
                    "-", field, "Record da inserire in %s: '%s'" % (model, value),
                    detail="usato in " + "; ".join(uses), value=value))

    def collect_child_emails(self):
        """Per ogni campo padre usato in [ref_child_email]: coppie (chiave del
        padre, e-mail) dei record figli di tutti i file che hanno sia il campo
        padre sia l'e-mail (es. CONTATTI: CODCONTO_REF + EMAILAZIENDA)"""
        for parent_field in set(self.cfg.ref_child_email.values()):
            pairs = self.child_emails.setdefault(parent_field, set())
            sources = self.child_files.setdefault(parent_field, [])
            for name, path in self.files.items():
                key = self.dirmap_of[name][0]
                cols = self.maps.columns.get(key, {})
                parent_cols = [c for c, mp in cols.items() if mp.field == parent_field]
                email_cols = [c for c, mp in cols.items() if mp.field == EMAIL_FIELD]
                if not parent_cols or not email_cols:
                    continue
                sources.append(name)
                header = None
                for _row, values, error in read_data_file(path, self.cfg.encoding):
                    if header is None:
                        header = [normalize_header(h) for h in values]
                        continue
                    if error or len(values) != len(header):
                        continue
                    row = dict(zip(header, (unquote(v) for v in values)))
                    for pcol in parent_cols:
                        parent = self.norm_key(row.get(pcol, ""))
                        if not parent:
                            continue
                        for ecol in email_cols:
                            for address in email_addresses(row.get(ecol, "")):
                                pairs.add((parent, address))

    def check_child_emails(self, parent_field, value, result):
        """Valori separati da ',' (come l'importatore): e-mail valide di
        contatti figli del partner della riga; l'importatore crea i mancanti"""
        row, _columns = self._row or ({}, {})
        key = self.dirmap_of[result.name][0]
        pk_cols = split_list(self.maps.dirmaps[key].pk) if key else []
        parent = next((row.get(c) for c in pk_cols if row.get(c)), "")
        bad, missing = [], []
        for tok in value.split(","):
            tok = tok.strip()
            if not tok:
                continue
            addresses = email_addresses(tok)
            if not addresses:
                bad.append(tok)
            elif parent and (self.norm_key(parent), addresses[0]) not in \
                    self.child_emails.get(parent_field, set()):
                missing.append(tok)
        if bad:
            return "E-mail %s non valida/e" % ", ".join("'%s'" % b for b in bad)
        if missing:
            return ("Contatto/i con e-mail %s non trovato/i tra i figli di '%s' "
                    "in %s: l'importazione li crea"
                    % (", ".join("'%s'" % m for m in missing), parent,
                       ", ".join(self.child_files.get(parent_field, [])) or "-"))
        return None

    def suggest_alternatives(self, model, field, lineno, result):
        """Testo con i record di anagrafica che hanno la stessa e-mail o la
        stessa partita IVA della riga; se nessuno, quelli con lo stesso
        dominio e-mail (non per i provider pubblici). None se la riga non ha
        né e-mail né partita IVA o se il campo non è in [ref_search]"""
        names = self.cfg.ref_search.get((model, field))
        if names is None or not self._row:
            return None
        row, columns = self._row
        addresses, vats = [], []
        for col, mapping in columns.items():
            if mapping.field == EMAIL_FIELD:
                addresses += email_addresses(row.get(col, ""))
            elif mapping.field in VAT_FIELDS + FISCALCODE_FIELDS:
                vat = normalize_vat(row.get(col, ""))
                if vat and vat not in vats:
                    vats.append(vat)
        if not addresses and not vats:
            return None
        this = (self.dirmap_of[result.name][0], lineno)

        def find(kind, keys):
            found = []
            for name in names:
                index = self.alt_index.get(name, {}).get(kind, {})
                for key in keys:
                    found += [r for r in index.get(key, [])
                              if r[:2] != this and r not in found]
            return found

        def describe(records):
            text = "; ".join(
                "%s riga %d: %s%s" % (n, row_, key or "-", " (%s)" % desc if desc else "")
                for n, row_, key, desc in records[:MAX_SUGGESTIONS])
            if len(records) > MAX_SUGGESTIONS:
                text += "; e altri %d" % (len(records) - MAX_SUGGESTIONS)
            return text

        def quoted(values):
            return ", ".join("'%s'" % v for v in values)

        parts = []
        records = find("email", addresses)
        if records:
            parts.append("stessa e-mail %s in: %s" % (quoted(addresses), describe(records)))
        records = find("vat", vats)
        if records:
            parts.append("stessa partita IVA %s in: %s" % (quoted(vats), describe(records)))
        if parts:
            return " | ".join(parts)
        searched = " e ".join(x for x, ok in (("e-mail", addresses), ("partita IVA", vats)) if ok)
        domains = sorted({a.split("@")[1] for a in addresses} - self.cfg.public_domains)
        if addresses and not domains:
            return ("nessun record con la stessa %s in %s; dominio di un provider "
                    "pubblico: ricerca per dominio non eseguita"
                    % (searched, ", ".join(names)))
        records = find("domain", domains)
        if records:
            return "nessun record con la stessa %s; stesso dominio %s in: %s" % (
                searched, quoted(domains), describe(records))
        # senza i valori della riga: le righe senza corrispondenze si raggruppano
        return "ricerca per %s%s: nessun record in %s" % (
            searched, " e per dominio" if domains else "", ", ".join(names))

    def fix_iban(self, name, path):
        """[iban]: riscrive il file ricalcolando le cifre di controllo "00"
        degli IBAN fittizi e togliendo le righe con IBAN non valido;
        l'originale è copiato in .bak (una volta sola). Le righe tolte sono
        segnalate con il numero di riga del file originale.
        Restituisce (IBAN corretti, righe tolte)"""
        key = self.dirmap_of[name][0]
        col = self.cfg.iban.get(name) or self.cfg.iban.get(key)
        if not col:
            return 0, 0
        issues = self.iban_issues.setdefault(name, [])
        with open(path, encoding=self.cfg.encoding, newline="") as fd:
            text = fd.read()
        eol = "\r\n" if "\r\n" in text else "\n"
        header = None
        kept = []
        fixed = removed = 0
        removed_rows = []
        for lineno, values, error in read_data_file(path, self.cfg.encoding):
            if header is None:
                header = [normalize_header(h) for h in values]
                if col not in header:
                    issues.append(Issue(
                        lineno, col, "[iban] colonna IBAN assente nel file"))
                    return 0, 0
                ix = header.index(col)
                kept.append("|".join(values))
                continue
            if error or len(values) != len(header):
                kept.append("|".join(values))  # segnalato dai controlli
                continue
            raw = values[ix]
            value = unquote(raw)
            if not value:
                kept.append("|".join(values))  # valore mancante: segnalato dai controlli
                continue
            iban = fix_iban_check_digits(value)
            msg = check_iban(iban)
            if msg:
                removed += 1
                removed_rows.append(Issue(
                    lineno, col, "%s: riga rimossa (riga del file .bak)" % msg,
                    value=value))
                continue
            if iban != value:
                if iban[2:4] != compact_iban(value)[2:4]:
                    fixed += 1
                quoted = RE_QUOTED.match(raw.strip())
                values[ix] = '"%s"' % iban if quoted else iban
            kept.append("|".join(values))
        if not (fixed or removed):
            return 0, 0
        backup = path + ".bak"
        if not os.path.exists(backup):
            shutil.copy2(path, backup)
        with open(path, "w", encoding=self.cfg.encoding.replace("-sig", ""),
                  newline="") as fd:
            fd.write(eol.join(kept) + eol)
        issues.append(Issue(
            "-", col, "[iban] %d IBAN con cifre di controllo 00 ricalcolati, "
            "%d righe con IBAN non valido rimosse (originale in %s)"
            % (fixed, removed, os.path.basename(backup))))
        issues.extend(removed_rows)
        return fixed, removed

    def check_file(self, name, path):
        key, by_prefix = self.dirmap_of[name]
        model = self.model_of(name)
        result = FileResult(name, path, model)
        result.issues.extend(self.iban_issues.get(name, []))
        unique = self.cfg.unique_only.get(name) or self.cfg.unique_only.get(key)
        if unique:
            return self.check_unique(path, unique, result)
        lookup = self.cfg.lookup_only.get(name) or self.cfg.lookup_only.get(key)
        if lookup:
            unique = self.cfg.unique.get(name) or self.cfg.unique.get(key) or []
            return self.check_lookup(path, lookup, unique, result)
        if name in self.nonconforming:
            result.issues.append(Issue(
                "-", "-",
                "Nome file non conforme a {NOME}_[0|1]_{YYYYmmdd}_{HHMM}.csv"))
        if not key:
            result.issues.append(Issue(
                "-", "-", "{NOME} '%s' assente in dirmap: file non importabile" % name))
        elif by_prefix:
            result.issues.append(Issue(
                "-", "-", "{NOME} '%s' assente in dirmap: usata la mappatura di %s"
                % (name, key)))
        if key:
            for msg in self.maps.problems.get(key, []):
                result.issues.append(Issue("-", "-", msg))
        columns = self.maps.columns.get(key, {}) if key else {}
        if key and not columns:
            result.issues.append(Issue("-", "-", "Nessuna colonna in mapper per %s" % key))
        unique = self.cfg.unique.get(name) or self.cfg.unique.get(key) or []
        requires = self.cfg.requires.get(name) or self.cfg.requires.get(key) or {}
        ref_code, ref_cols = (self.cfg.row_ref.get(name) or self.cfg.row_ref.get(key)
                              or ("", []))
        seen = {}  # combinazione normalizzata di [unique] -> prima riga
        header = None
        for lineno, values, error in read_data_file(path, self.cfg.encoding):
            if header is None:
                header = [normalize_header(h) for h in values]
                self.check_header(lineno, header, columns, result)
                if key:
                    self.check_pk(lineno, header, key, result)
                for col in unique + [c for c in requires if c not in unique]:
                    if col not in header:
                        result.issues.append(Issue(
                            lineno, col, "Colonna di [%s] assente nel file"
                            % ("unique" if col in unique else "requires")))
                continue
            result.records += 1
            if error:
                issues = [Issue(lineno, "-", error)]
            else:
                issues = self.check_row(lineno, header, values, columns, result)
                if len(values) == len(header):
                    row = dict(zip(header, (unquote(v) for v in values)))
                    issues += self.check_requires(lineno, row, requires)
                    issues += self.check_unique_cols(lineno, row, unique, seen)
                    if row.get(ref_code):
                        for issue in issues:
                            issue.ref = row[ref_code]
                            issue.pinned = issue.field in ref_cols
            if issues:
                result.bad_records += 1
                result.issues.extend(issues)
        if header is None:
            result.issues.append(Issue("-", "-", "File vuoto"))
        return result

    def check_unique(self, path, unique, result):
        """[unique_only]: nessun valore ripetuto in ciascuna delle colonne
        indicate; il file non viene controllato in altro modo"""
        header = None
        seen = {}  # colonna -> {valore normalizzato: prima riga}
        for lineno, values, error in read_data_file(path, self.cfg.encoding):
            if header is None:
                header = [normalize_header(h) for h in values]
                for col in unique:
                    if col not in header:
                        result.issues.append(Issue(
                            lineno, col, "Colonna da controllare per duplicati "
                            "assente nel file"))
                continue
            result.records += 1
            if error:
                issues = [Issue(lineno, "-", error)]
            elif len(values) != len(header):
                issues = [Issue(
                    lineno, "-",
                    "Numero colonne %d diverso dall'intestazione (%d)"
                    % (len(values), len(header)))]
            else:
                issues = []
                for col in unique:
                    if col not in header:
                        continue
                    value = unquote(values[header.index(col)])
                    if not value:
                        continue
                    first = seen.setdefault(col, {}).setdefault(
                        self.norm_key(value), lineno)
                    if first != lineno:
                        issues.append(Issue(
                            lineno, col, "Valore '%s' duplicato (già alla riga %s)"
                            % (value, first), value=value))
            if issues:
                result.bad_records += 1
                result.issues.extend(issues)
        if header is None:
            result.issues.append(Issue("-", "-", "File vuoto"))
        return result

    def lookup_values(self, source, target):
        """Valori validi di una regola di [lookup_only]: colonna target del
        file source, oppure chiavi del modello source (target None).
        Restituisce (insieme normalizzato, descrizione) o (None, errore)"""
        if target is None:
            if source not in self._models_with_data:
                return None, "nessun file dati fornisce le chiavi di %s" % source
            files = ", ".join(sorted(self.ref_files.get(source, [])))
            return self.ref_keys.get(source, set()), "%s (%s)" % (files, source)
        path = self.files.get(source) or self.lookup_files.get(source)
        if not path:
            path = os.path.join(self.cfg.data_dir, source)
            if not os.path.isfile(path):
                path = next(
                    (os.path.join(self.cfg.data_dir, fn)
                     for fn in sorted(os.listdir(self.cfg.data_dir))
                     if fn.lower() == source.lower()), None)
        if not path:
            return None, "file %s non trovato in %s" % (source, self.cfg.data_dir)
        values = read_column(path, target, self.cfg.encoding)
        if values is None:
            return None, "colonna '%s' assente in %s" % (target, os.path.basename(path))
        return ({self.norm_key(v) for v in values},
                "'%s' di %s" % (target, os.path.basename(path)))

    def check_lookup(self, path, rules, unique, result):
        """[lookup_only]: i valori delle colonne indicate devono comparire
        nella colonna di un altro file o tra le chiavi di un modello; il file
        non viene controllato in altro modo, salvo le colonne di [unique]"""
        header = None
        seen = {}  # combinazione normalizzata di [unique] -> prima riga
        checks = []  # (indice colonna, colonna, valori validi, descrizione)
        for lineno, values, error in read_data_file(path, self.cfg.encoding):
            if header is None:
                header = [normalize_header(h) for h in values]
                for col, source, target in rules:
                    valid, where = self.lookup_values(source, target)
                    if valid is None:
                        result.issues.append(Issue(
                            lineno, col, "Valori non verificabili: %s" % where))
                    elif col not in header:
                        result.issues.append(Issue(
                            lineno, col, "Colonna da verificare assente nel file"))
                    else:
                        checks.append((header.index(col), col, valid, where))
                for col in unique:
                    if col not in header:
                        result.issues.append(Issue(
                            lineno, col, "Colonna di [unique] assente nel file"))
                continue
            result.records += 1
            if error:
                issues = [Issue(lineno, "-", error)]
            elif len(values) != len(header):
                issues = [Issue(
                    lineno, "-",
                    "Numero colonne %d diverso dall'intestazione (%d)"
                    % (len(values), len(header)))]
            else:
                issues = []
                for ix, col, valid, where in checks:
                    value = unquote(values[ix])
                    if not value:
                        issues.append(Issue(lineno, col, "Valore mancante"))
                    elif self.norm_key(value) not in valid:
                        issues.append(Issue(
                            lineno, col, "Valore '%s' non trovato in %s"
                            % (value, where), value=value))
                row = dict(zip(header, (unquote(v) for v in values)))
                issues += self.check_unique_cols(lineno, row, unique, seen)
            if issues:
                result.bad_records += 1
                result.issues.extend(issues)
        if header is None:
            result.issues.append(Issue("-", "-", "File vuoto"))
        return result

    def purge_file(self, result):
        """[purge] (temporaneo): solo nella data indicata riscrive il file
        senza i record con errori; l'originale è copiato in .bak (una volta
        sola, per non perderlo con una seconda esecuzione)"""
        key = self.dirmap_of[result.name][0]
        day = self.cfg.purge.get(result.name) or self.cfg.purge.get(key)
        if day != date.today().isoformat():
            return 0
        bad = {i.row for i in result.issues if isinstance(i.row, int)}
        with open(result.path, encoding=self.cfg.encoding, newline="") as fd:
            text = fd.read()
        eol = "\r\n" if "\r\n" in text else "\n"
        kept = []
        removed = 0
        for lineno, values, _error in read_data_file(result.path, self.cfg.encoding):
            if kept and lineno in bad:  # mai l'intestazione
                removed += 1
            else:
                kept.append("|".join(values))
        if not removed:
            return 0
        backup = result.path + ".bak"
        if not os.path.exists(backup):
            shutil.copy2(result.path, backup)
        with open(result.path, "w", encoding=self.cfg.encoding.replace("-sig", ""),
                  newline="") as fd:
            fd.write(eol.join(kept) + eol)
        result.issues.insert(0, Issue(
            "-", "-", "[purge] %d record con errori rimossi dal file "
            "(originale in %s)" % (removed, os.path.basename(backup))))
        return removed

    def check_unique_cols(self, lineno, row, unique, seen):
        """[unique]: la combinazione dei valori delle colonne non deve
        ripetersi in altre righe del file (righe con tutte vuote escluse)"""
        if not unique or not all(col in row for col in unique):
            return []
        values = [row[col] for col in unique]
        if not any(values):
            return []
        first = seen.setdefault(
            tuple(self.norm_key(v) for v in values), lineno)
        if first == lineno:
            return []
        combo = ", ".join("%s='%s'" % (c, v) for c, v in zip(unique, values))
        return [Issue(
            lineno, "+".join(unique),
            "Combinazione duplicata (già alla riga %s): %s" % (first, combo))]

    @staticmethod
    def check_requires(lineno, row, requires):
        """[requires]: se la colonna ha un valore, anche le colonne
        indicate devono averlo"""
        issues = []
        for col, needed in requires.items():
            if not row.get(col):
                continue
            for other in needed:
                if other in row and not row[other]:
                    issues.append(Issue(
                        lineno, other, "Valore mancante: obbligatorio se %s "
                        "è impostato ('%s')" % (col, row[col]), value=row[col]))
        return issues

    def check_header(self, lineno, header, columns, result):
        seen = set()
        for ix, col in enumerate(header):
            if not col:
                result.issues.append(Issue(
                    lineno, "#%d" % (ix + 1), "Intestazione colonna vuota"))
            elif columns and col not in columns:
                result.issues.append(Issue(
                    lineno, col, "Colonna assente in mapper: non importata"))
            if col in seen:
                result.issues.append(Issue(
                    lineno, col, "Colonna duplicata nell'intestazione"))
            seen.add(col)
        for col, mapping in columns.items():
            if mapping.needs_value and col not in header:
                result.issues.append(Issue(
                    lineno, col, "Colonna obbligatoria mancante (%s)"
                    % mapping.required))

    def check_pk(self, lineno, header, key, result):
        """La colonna chiave di dirmap (counterpart_pk) deve esistere nel file"""
        dirmap = self.maps.dirmaps[key]
        pks = split_list(dirmap.pk)
        if not pks:
            result.issues.append(Issue(
                lineno, "-", "dirmap: counterpart_pk non impostato per %s" % key))
        for pk in pks:
            if pk not in header:
                result.issues.append(Issue(
                    lineno, pk,
                    "dirmap: colonna chiave (counterpart_pk) assente nel file; "
                    "i riferimenti a questo file non sono verificabili (%s)"
                    % dirmap.model))

    def check_row(self, lineno, header, values, columns, result):
        if len(values) != len(header):
            return [Issue(
                lineno, "-",
                "Numero colonne %d diverso dall'intestazione (%d)"
                % (len(values), len(header)))]
        issues = []
        country = None
        self._row = None
        if self.cfg.ref_search or self.cfg.ref_child_email:
            self._row = (dict(zip(header, (unquote(v) for v in values))), columns)
        for col, raw in zip(header, values):
            mapping = columns.get(col)
            if not mapping or not mapping.field:
                continue
            value = unquote(raw)
            if mapping.field in VAT_FIELDS and value and not mapping.apply4:
                if country is None:
                    country = self.row_country(header, values, columns)
                msg = check_vat(value, country)
                if msg:
                    issues.append(Issue(lineno, col, msg, value=value))
                continue
            if not value:
                if mapping.needs_value:
                    issues.append(Issue(
                        lineno, col, "Valore obbligatorio mancante: %s (%s)"
                        % (mapping.field, mapping.required)))
                continue
            if mapping.apply4:
                continue  # valore trasformato in importazione
            msg = self.check_value(mapping, value, lineno, result)
            if isinstance(msg, tuple):  # (messaggio, suggerimenti di [ref_search])
                issues.append(Issue(lineno, col, *msg, value=value, searched=True))
            elif msg:
                issues.append(Issue(lineno, col, msg, value=value))
        return issues

    @staticmethod
    def row_country(header, values, columns):
        """Codice ISO del paese della riga: colonna mappata su country_id,
        altrimenti il suo default"""
        for col, mapping in columns.items():
            if mapping.field != COUNTRY_FIELD:
                continue
            value = ""
            if col in header:
                value = unquote(values[header.index(col)])
            return (value or mapping.default or "").strip().upper()
        return ""

    def check_value(self, mapping, value, lineno, result):
        ttype = mapping.ttype
        if ttype in REF_TYPES:
            msg = self.check_ref_suffix(mapping, value, result)
            if msg:
                return msg  # la causa è il suffisso: niente ricerca del riferimento
            parent_field = self.option(self.cfg.ref_child_email, result, mapping.field)
            if parent_field:  # riferimenti per e-mail ai contatti figli
                return self.check_child_emails(parent_field, value, result)
            refs = split_refs(value) if ttype in X2MANY_TYPES else [value]
            return self.check_refs(mapping, refs, lineno, result)
        if ttype in INTEGER_TYPES and not RE_INTEGER.match(value):
            return "Valore '%s' non è un numero intero" % value
        if ttype in FLOAT_TYPES and not RE_FLOAT.match(value):
            return ("Valore '%s' non è un numero in formato USA "
                    "(punto decimale, senza separatore migliaia)" % value)
        if ttype == "selection" and mapping.selection:
            return check_selection(mapping.selection, value)
        if ttype in EMAIL_TYPES and RE_EMAIL_FIELD.search(mapping.field):
            return check_email(value)
        if mapping.field in FISCALCODE_FIELDS:
            return check_fiscalcode(value)
        if ttype in EMAIL_TYPES and RE_WEBSITE_FIELD.search(mapping.field):
            return check_website(value)
        if ttype == "date" and not valid_date(value):
            return "Data '%s' non valida o non in formato ISO YYYY-mm-dd" % value
        if ttype == "datetime" and not valid_datetime(value):
            return ("Data/ora '%s' non valida o non in formato ISO "
                    "YYYY-mm-dd HH:MM:SS" % value)
        return None

    def check_ref_suffix(self, mapping, value, result):
        """[ref_suffix]: ogni riferimento al modello deve terminare con il
        suffisso indicato (es. le categorie prodotto foglia con '.')"""
        suffix = self.cfg.ref_suffix.get(mapping.relation)
        # esclusione per nome esatto del file (CONTATTI non esclude CONTATTI_SUPPLIERS)
        if not suffix or {result.name, self.dirmap_of[result.name][0]} & \
                self.cfg.ref_suffix_exclude:
            return None
        refs = split_refs(value) if mapping.ttype in X2MANY_TYPES else [value]
        bad = [r for r in refs if not r.endswith(suffix)]
        if bad:
            return "Riferimento/i %s a %s senza '%s' finale" % (
                ", ".join("'%s'" % r for r in bad), mapping.relation, suffix)
        return None

    def check_refs(self, mapping, refs, lineno, result):
        relation = mapping.relation
        if not relation or relation in self.cfg.external_models:
            return None
        if relation not in self._models_with_data:
            tag = (mapping.column, relation)
            if tag not in result.warned:
                result.warned.add(tag)
                result.issues.append(Issue(
                    lineno, mapping.column,
                    "Riferimenti non verificabili: nessun file dati fornisce "
                    "le chiavi (%s; aggiungerlo a [external] se già in Odoo)"
                    % relation))
            return None
        only = self.option(self.cfg.ref_files, result, mapping.field)
        if only:  # solo le chiavi dei file indicati in [ref_files]
            cache_key = tuple(only)
            if cache_key not in self._only_keys:
                self._only_keys[cache_key] = set().union(
                    *(self.name_keys.get(n, set()) for n in only))
            keys = self._only_keys[cache_key]
            sources = only
        else:
            keys = self.ref_keys.get(relation, set())
            sources = sorted(self.ref_files.get(relation, []))
        missing = [r for r in refs if self.norm_key(r) not in keys]
        allowed = self.option(self.cfg.ref_values, result, mapping.field)
        if missing and allowed:  # valori sempre validi di [ref_values]
            allowed = {self.norm_key(x) for x in allowed}
            missing = [r for r in missing if self.norm_key(r) not in allowed]
        target = self.option(self.cfg.ref_insert, result, mapping.field)
        if missing and target:
            # non è un errore di questo file: record da inserire nel file collegato
            self._insert_model[target] = relation
            values = self.to_insert.setdefault(target, {})
            for ref in missing:
                values.setdefault(ref, {}).setdefault(result.name, []).append(lineno)
            return None
        if missing:
            msg = "Riferimento/i %s non trovato/i in %s (%s)" % (
                ", ".join("'%s'" % r for r in missing), ", ".join(sources), relation)
            hint = self.suggest_alternatives(result.model, mapping.field, lineno, result)
            return (msg, hint) if hint else msg
        return None


def parse_selection(value):
    """'{"chiave": "etichetta"}' -> dict; None se assente o non leggibile"""
    if not value:
        return None
    try:
        data = json.loads(value)
    except ValueError:
        return None
    return data if isinstance(data, dict) and data else None


def check_selection(selection, value):
    """Come l'import di Odoo: chiave o etichetta, senza distinguere
    maiuscole e minuscole"""
    low = value.lower()
    for key, label in selection.items():
        if low == str(key).lower() or (label and low == str(label).lower()):
            return None
    return "Valore '%s' non ammesso (valori validi: %s)" % (
        value, ", ".join("%s=%s" % (k, l) if l and l != k else str(k)
                         for k, l in selection.items() if k != ""))


def valid_it_vat(number):
    """Partita IVA italiana: 11 cifre, codice ufficio valido e cifra di
    controllo (stesso algoritmo di stdnum.it.iva)"""
    if not re.fullmatch(r"\d{11}", number) or int(number[:7]) == 0:
        return False
    office = int(number[7:10])
    if not (1 <= office <= 100 or office in (120, 121, 888, 999)):
        return False
    total = 0
    for ix, digit in enumerate(number[:10]):
        n = int(digit) * (2 if ix % 2 else 1)
        total += n - 9 if n > 9 else n
    return (10 - total % 10) % 10 == int(number[10])


def valid_person_fiscalcode(code):
    """Codice fiscale di persona fisica (16 caratteri): formato, mese,
    giorno (1-31, +40 per le donne; più severo di stdnum, che accetta il
    giorno modulo 40) e carattere di controllo"""
    if not RE_FISCALCODE.match(code):
        return False
    digits = "".join(FC_OMOCODE.get(c, c) if ix in (6, 7, 9, 10, 12, 13, 14) else c
                     for ix, c in enumerate(code))
    day = int(digits[9:11])
    day = day - 40 if day > 40 else day
    month = FC_MONTHS.index(code[8]) + 1
    leap = int(digits[6:8]) % 4 == 0  # 19yy o 20yy: bisestile se yy % 4 == 0
    days = (29 if leap else 28) if month == 2 else 30 if month in (4, 6, 9, 11) else 31
    if not 1 <= day <= days:
        return False
    total = sum(FC_ODD[c] if ix % 2 == 0 else
                (int(c) if c.isdigit() else ord(c) - ord("A"))
                for ix, c in enumerate(code[:15]))
    return chr(ord("A") + total % 26) == code[15]


def check_fiscalcode(value):
    """Codice fiscale: 16 caratteri (persona fisica) oppure partita IVA
    italiana di 11 cifre (anche con prefisso IT, che Odoo toglie)"""
    code = re.sub(r"[\s.\-]", "", value).upper()
    if len(code) == 16:
        if valid_person_fiscalcode(code):
            return None
    elif valid_it_vat(code[2:] if code.startswith("IT") else code):
        return None
    return ("Codice fiscale '%s' non valido (16 caratteri o partita IVA "
            "di 11 cifre)" % value)


def check_vat(value, country):
    """Come base_vat di Odoo: il prefisso (es. IT, DE) indica il paese,
    altrimenti vale il paese del partner"""
    vat = re.sub(r"[\s.\-/]", "", value).upper()
    match = RE_VAT_PREFIX.match(vat)
    if match and (match.group(1) == country or not match.group(2)[:1].isalpha()):
        country, number = match.group(1), match.group(2)
    else:
        number = vat
    if not country:
        return None  # senza paese Odoo non la verifica
    if country == "IT":
        if valid_it_vat(number):
            return None
        return "Partita IVA '%s' non valida (IT: 11 cifre con cifra di controllo)" % value
    if get_cc_module is None:
        return None  # python-stdnum non installato: estere non verificate
    cc = country.lower().replace("el", "gr")
    modules = [get_cc_module(cc, "vat")]
    if cc == "ch":  # Odoo accetta anche il solo UID (CHE-123.456.789)
        modules.append(get_cc_module(cc, "uid"))
    modules = [mod for mod in modules if mod is not None]
    if not modules or any(mod.is_valid(v) for mod in modules
                          for v in (country + number, number, vat)):
        return None
    return "Partita IVA '%s' non valida per il paese %s" % (value, country)


def compact_iban(value):
    """IBAN senza spazi, in maiuscolo"""
    return re.sub(r"\s", "", value).upper()


def iban_remainder(iban):
    """Resto modulo 97 dell'IBAN (paese e cifre di controllo in coda,
    lettere convertite in numeri: A=10 ... Z=35); 1 se l'IBAN è valido"""
    text = iban[4:] + iban[:4]
    return int("".join(str(int(c, 36)) for c in text)) % 97


def iban_wellformed(iban):
    """Formato dell'IBAN compattato: paese noto, lunghezza del paese,
    caratteri alfanumerici e, per IT/SM, struttura del BBAN"""
    if not RE_IBAN.match(iban) or IBAN_LENGTHS.get(iban[:2]) != len(iban):
        return False
    bban = RE_IBAN_BBAN.get(iban[:2])
    return not bban or bool(bban.match(iban[4:]))


def fix_iban_check_digits(value):
    """IBAN fittizio con cifre di controllo "00" (es. IT00A...): restituisce
    l'IBAN compattato con le cifre ricalcolate (98 - resto modulo 97);
    gli altri valori sono solo compattati. "00" non è mai valido"""
    iban = compact_iban(value)
    if iban[2:4] == "00" and iban_wellformed(iban):
        iban = "%s%02d%s" % (iban[:2], 98 - iban_remainder(iban), iban[4:])
    return iban


def check_iban(value):
    """IBAN: formato del paese e cifra di controllo modulo 97 (il CIN
    italiano interno al BBAN non viene verificato)"""
    iban = compact_iban(value)
    if not iban_wellformed(iban):
        return "IBAN '%s' non valido (formato o lunghezza)" % value
    if iban_remainder(iban) != 1:
        return "IBAN '%s' non valido (cifre di controllo errate)" % value
    return None


def check_website(value):
    """Un solo indirizzo web; spazi e a capo iniziali/finali sono ammessi.
    Con www. servono almeno tre parti (www.ditta.it, non www.ditta)"""
    url = value.strip()
    if RE_WEBSITE.match(url):
        host = re.sub(r"^https?://", "", url, flags=re.IGNORECASE)
        host = re.split(r"[:/?#]", host, maxsplit=1)[0].lower()
        if not host.startswith("www.") or host.count(".") >= 2:
            return None
    return "Sito web '%s' non valido" % value


def split_emails(value):
    """Separa su ',' o ';' fuori da virgolette e da <...>"""
    parts, buf, quoted, angle = [], [], False, False
    for ch in value:
        if ch == '"' and not angle:
            quoted = not quoted
        elif ch == "<" and not quoted:
            angle = True
        elif ch == ">" and not quoted:
            angle = False
        elif ch in ",;" and not quoted and not angle:
            parts.append("".join(buf))
            buf = []
            continue
        buf.append(ch)
    parts.append("".join(buf))
    return parts


def normalize_vat(value):
    """Partita IVA confrontabile: senza spazi e separatori, maiuscola, senza
    prefisso IT per le italiane; un codice fiscale di 16 caratteri (persona
    fisica) non è una partita IVA e viene scartato"""
    vat = re.sub(r"[^A-Za-z0-9]", "", value or "").upper()
    if re.fullmatch(r"IT\d{11}", vat):
        vat = vat[2:]
    if len(vat) == 16 and not vat.isdigit():
        return ""
    if len(vat) < 8 or set(vat) <= {"0", "9"}:
        return ""  # troppo corta o segnaposto (00000000000, 99999999999)
    return vat


def email_addresses(value):
    """Indirizzi e-mail validi contenuti nel valore, in minuscolo"""
    result = []
    for part in split_emails(value or ""):
        part = part.strip()
        match = RE_EMAIL_FORMATTED.match(part)
        address = (match.group(1) if match else part).strip().lower()
        if RE_EMAIL.match(address) and address not in result:
            result.append(address)
    return result


def check_email(value):
    """Uno o più indirizzi separati da ',' o ';', anche nella forma
    'Nome <indirizzo>'; spazi e a capo iniziali/finali sono ammessi"""
    bad = []
    for part in split_emails(value):
        part = part.strip()
        if not part:
            continue
        match = RE_EMAIL_FORMATTED.match(part)
        address = match.group(1).strip() if match else part
        if not RE_EMAIL.match(address):
            bad.append(part)
    if not bad and not value.strip(" \t\r\n\xa0,;"):
        bad.append(value)
    if bad:
        return "E-mail %s non valida/e" % ", ".join("'%s'" % b for b in bad)
    return None


def valid_date(value):
    if not RE_DATE.match(value):
        return False
    try:
        date.fromisoformat(value)
    except ValueError:
        return False
    return True


def valid_datetime(value):
    if not RE_DATETIME.match(value):
        return False
    value = value.replace("Z", "+00:00")
    if re.search(r"[+-]\d{4}$", value):
        value = value[:-2] + ":" + value[-2:]
    # Python 3.10: fromisoformat accetta solo 3 o 6 cifre decimali
    value = re.sub(r"\.(\d{1,6})", lambda m: "." + m.group(1).ljust(6, "0"), value)
    try:
        datetime.fromisoformat(value)
    except ValueError:
        return False
    return True


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------
def summary_rows(run_time, results):
    tot = sum(r.records for r in results)
    bad = sum(r.bad_records for r in results)
    return [
        ("Data e ora del controllo", run_time.strftime("%Y-%m-%d %H:%M:%S")),
        ("File esaminati", len(results)),
        ("Record esaminati", tot),
        ("Record errati", bad),
        ("Percentuale record validi", fmt_pct(pct(tot, bad))),
    ]


def md_cell(value):
    return str(value).replace("|", "\\|").replace("\n", " ")


def md_anchor(title):
    """Ancora di un titolo come la calcola GitHub / GitLab"""
    return re.sub(r"[^\w\- ]", "", title.strip().lower()).replace(" ", "-")


def write_md(path, run_time, results, skipped, matches=()):
    out = ["# Controllo file CSV per Odoo 18.0", "", "## Indice", ""]
    titles = (["Riepilogo file"] + [t for t, _h, _r in matches]
              + [os.path.basename(r.path) for r in results])
    if skipped:
        titles.append("File non esaminati")
    out += ["- [%s](#%s)" % (md_cell(t), md_anchor(t)) for t in titles]
    out.append("")
    out += ["| Voce | Valore |", "|---|---|"]
    out += ["| %s | %s |" % (k, md_cell(v)) for k, v in summary_rows(run_time, results)]
    out += ["", "## Riepilogo file", "",
            "| File | Modello | Record | Errati | Validi |", "|---|---|--:|--:|--:|"]
    for r in results:
        out.append("| %s | %s | %d | %d | %s |" % (
            md_cell(os.path.basename(r.path)), r.model or "-",
            r.records, r.bad_records, fmt_pct(r.valid_pct)))
    for title, header, rows in matches:
        out += ["", "## %s" % title, "",
                "Record con la stessa partita IVA: %d" % len(rows), ""]
        if rows:
            out += ["| %s |" % " | ".join(md_cell(c) for c in header),
                    "|%s" % ("---|" * len(header))]
            out += ["| %s |" % " | ".join(md_cell(c) for c in row) for row in rows]
    for r in results:
        out += ["", "## %s" % os.path.basename(r.path), "",
                "- Nome completo: `%s`" % r.path,
                "- Modello: %s" % (r.model or "-"),
                "- Record esaminati: %d" % r.records,
                "- Record errati: %d" % r.bad_records,
                "- Percentuale record validi: %s" % fmt_pct(r.valid_pct), ""]
        if r.issues:
            out += ["| Riga | Campo | Errore |", "|--:|---|---|"]
            out += ["| %s | %s | %s |" % (md_cell(i.label), md_cell(i.field),
                                          md_cell(i.message))
                    for i in r.issues]
        else:
            out.append("Nessuna segnalazione.")
    if skipped:
        out += ["", "## File non esaminati", "", "| File | Motivo |", "|---|---|"]
        out += ["| %s | %s |" % (md_cell(f), md_cell(m)) for f, m in skipped]
    with open(path, "w", encoding="utf-8") as fd:
        fd.write("\n".join(out) + "\n")


def report_sheets(run_time, results, skipped, matches=()):
    """Contenuto del report come fogli: [(nome, [righe])]"""
    summary = [["Voce", "Valore"]] + [list(x) for x in summary_rows(run_time, results)]
    summary += [[], ["File", "Nome completo", "Modello", "Record esaminati",
                     "Record errati", "% record validi"]]
    for r in results:
        summary.append([os.path.basename(r.path), r.path, r.model or "-",
                        r.records, r.bad_records,
                        "" if r.valid_pct is None else r.valid_pct])
    errors = [["File", "Riga", "Codice", "Campo", "Errore"]]
    for r in results:
        for i in r.issues:
            errors.append([os.path.basename(r.path), i.row, i.ref, i.field, i.message])
    # nome del foglio: al massimo 31 caratteri
    sheets = ([("Riepilogo", summary)]
              + [(title.replace("Partita IVA in ", "IVA ").replace(" e ", "-")[:31],
                  [header] + rows) for title, header, rows in matches]
              + [("Errori", errors)])
    if skipped:
        sheets.append(("Non esaminati", [["File", "Motivo"]] + [list(x) for x in skipped]))
    return sheets


def write_xlsx(path, run_time, results, skipped, matches=()):
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        raise SystemExit("Per il report .xlsx serve openpyxl (pip install openpyxl); "
                         "in alternativa usare .ods o .md")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in report_sheets(run_time, results, skipped, matches):
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 80)
    wb.save(path)


def write_ods(path, run_time, results, skipped, matches=()):
    """Scrive un .ods (OpenDocument) senza dipendenze esterne"""
    def cell(value):
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return ('<table:table-cell office:value-type="float" office:value="%s">'
                    '<text:p>%s</text:p></table:table-cell>' % (value, value))
        return ('<table:table-cell office:value-type="string"><text:p>%s</text:p>'
                '</table:table-cell>' % escape(str(value)))

    tables = []
    for title, rows in report_sheets(run_time, results, skipped, matches):
        xml = ['<table:table table:name="%s">' % escape(title)]
        for row in rows:
            xml.append("<table:table-row>%s</table:table-row>"
                       % ("".join(cell(v) for v in row) or "<table:table-cell/>"))
        xml.append("</table:table>")
        tables.append("".join(xml))
    ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')
    content = ('<?xml version="1.0" encoding="UTF-8"?>'
               '<office:document-content %s office:version="1.2"><office:body>'
               '<office:spreadsheet>%s</office:spreadsheet></office:body>'
               '</office:document-content>' % (ns, "".join(tables)))
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:'
        'xmlns:manifest:1.0" manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/" manifest:media-type='
        '"application/vnd.oasis.opendocument.spreadsheet"/>'
        '<manifest:file-entry manifest:full-path="content.xml" '
        'manifest:media-type="text/xml"/></manifest:manifest>')
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet",
                    compress_type=zipfile.ZIP_STORED)
        zf.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("META-INF/manifest.xml", manifest,
                    compress_type=zipfile.ZIP_DEFLATED)


RE_QUOTED_VALUE = re.compile(r"'[^']*'")
RE_ELLIPSIS_LIST = re.compile(r"'…'(?:, '…')+")
PDF_SAMPLE_VALUES = 6
PDF_ROWS_PER_CELL = 60

PDF_CSS = """
body { font-family: "DejaVu Sans", Arial, sans-serif; font-size: 9pt; color: #222; }
h1 { font-size: 17pt; margin: 0 0 4pt 0; }
h2 { font-size: 12pt; margin: 18pt 0 4pt 0; padding-bottom: 2pt;
     border-bottom: 1.5pt solid #2b5797; page-break-after: avoid; }
.sub { color: #666; margin-bottom: 10pt; }
.meta { color: #555; margin: 0 0 6pt 0; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 8pt 0; }
thead { display: table-header-group; }
tr { page-break-inside: avoid; }
th { background: #2b5797; color: #fff; text-align: left; padding: 3pt 5pt; }
td { border-bottom: 0.5pt solid #ccc; padding: 3pt 5pt; vertical-align: top; }
td.n, th.n { text-align: right; white-space: nowrap; }
.ok { color: #1e7b34; font-weight: bold; }
.warn { color: #b36b00; font-weight: bold; }
.bad { color: #b3261e; font-weight: bold; }
.small { font-size: 8pt; color: #555; word-break: break-all; }
.rows { font-size: 8pt; color: #555; }
.nb { white-space: pre; }
.hint { font-size: 8pt; font-style: italic; color: #1f4e79; margin-top: 2pt; }
tr.grp td { background: #e8eef7; border-top: 1pt solid #2b5797; padding: 4pt 5pt;
            page-break-after: avoid; }
tr.grp .cnt { float: right; font-weight: bold; color: #2b5797; }
tr.one td { border-top: 1pt solid #2b5797; }
th.file { background: #1f3f6e; font-size: 9pt; }
.val { font-family: "DejaVu Sans Mono", monospace; white-space: pre-wrap; }
table.issues { table-layout: fixed; }
.kpi td { border: none; padding: 2pt 12pt 2pt 0; }
.kpi td.v { font-size: 12pt; font-weight: bold; }
"""


def pct_class(value):
    if value is None:
        return ""
    return "ok" if value >= 99.5 else "warn" if value >= 90 else "bad"


def group_issues(issues):
    """Segnalazioni raggruppate per campo ed errore (valori tra apici
    sostituiti da '…'), ordinate per campo e messaggio:
    [(campo, messaggio, conteggio, voci, messaggio completo se un solo record,
    ricerca alternativa eseguita)];
    voci = [(righe, valore, dettaglio)] con le righe dello stesso valore e
    dettaglio unite ('3-5, 9')"""
    groups = {}
    for issue in issues:
        # uno o più valori ('…', '…') danno lo stesso messaggio
        template = RE_ELLIPSIS_LIST.sub("'…'", RE_QUOTED_VALUE.sub("'…'", issue.base))
        groups.setdefault((str(issue.field), template), []).append(issue)
    # stesso errore sulla stessa riga per più campi (es. intestazione): una voce
    single = {}
    for (field, template), items in list(groups.items()):
        if len(items) == 1:
            issue = items[0]
            single.setdefault((template, issue.label, issue.detail),
                              [[], issue.base])[0].append(field)
            del groups[(field, template)]
    result = []
    for (field, template), items in groups.items():
        entries = {}
        for issue in items:
            quoted = [q[1:-1] for q in RE_QUOTED_VALUE.findall(issue.base)]
            value = ", ".join(quoted) if quoted else issue.value
            # nel PDF solo i suggerimenti che indicano dei record
            detail = ("" if "nessun record" in issue.detail and " riga " not in issue.detail
                      else issue.detail)
            entries.setdefault((value, detail), []).append(issue)
        voci = sorted(((compress_rows(rows), value, detail)
                       for (value, detail), rows in entries.items()),
                      key=lambda v: first_row(v[0]))
        result.append((field, template, len(items), voci, None,
                       any(i.searched for i in items)))
    for (template, row, detail), (fields, base) in single.items():
        result.append((", ".join(fields), template, 1, [(str(row), None, detail)], base,
                       bool(detail)))
    return sorted(result, key=lambda g: (g[0] != "-", g[0], g[1]))


def first_row(rows):
    match = re.match(r"\d+", rows)
    return int(match.group()) if match else 0


def compress_rows(items):
    """righe (numeri o Issue) [3, 4, 5, 9, '-'] -> '3-5, 9, -' (trattino non
    divisibile); le righe con il codice del record ([row_ref]) restano
    singole: '7 (C  12)'; il codice compare per le colonne di [row_ref] e
    quando la voce ha una sola riga"""
    issues = [i for i in items if isinstance(i, Issue)]
    alone = len({i.row for i in issues}) == 1
    labels = {i.row: i.label for i in issues
              if i.ref and (i.pinned or alone) and isinstance(i.row, int)}
    rows = [i.row for i in issues if i.row not in labels]
    rows += [r for r in items if not isinstance(r, Issue)]
    nums = sorted({r for r in rows if isinstance(r, int)})
    parts = []
    ix = 0
    while ix < len(nums):
        end = ix
        while end + 1 < len(nums) and nums[end + 1] == nums[end] + 1:
            end += 1
        parts.append(str(nums[ix]) if end == ix else "%d\u2011%d" % (nums[ix], nums[end]))
        ix = end + 1
    parts += [labels[r] for r in labels]
    parts.sort(key=first_row)
    others = sorted({str(r) for r in rows if not isinstance(r, int)})
    return ", ".join(parts + others)


def pdf_rows(rows):
    """Elenco righe per il PDF: ogni 'riga (codice)' non va a capo e
    conserva gli spazi interni del codice"""
    return ", ".join("<span class='nb'>%s</span>" % escape(r) if " (" in r
                     else escape(r) for r in rows.split(", "))


def pdf_hint(detail):
    return "<div class='hint'>%s</div>" % escape(detail) if detail else ""


def write_pdf(path, run_time, results, skipped, matches=()):
    """Scrive un PDF tramite wkhtmltopdf (lo stesso usato da Odoo)"""
    import shutil
    import subprocess
    import tempfile

    wkhtmltopdf = shutil.which("wkhtmltopdf")
    if not wkhtmltopdf:
        raise SystemExit("Per il report .pdf serve wkhtmltopdf; "
                         "in alternativa usare .md, .xlsx o .ods")
    h = escape
    tot = sum(r.records for r in results)
    bad = sum(r.bad_records for r in results)
    total_pct = pct(tot, bad)
    out = [
        "<html><head><meta charset='utf-8'><style>%s</style></head><body>" % PDF_CSS,
        "<h1>Controllo file CSV per Odoo 18.0</h1>",
        "<div class='sub'>Controllo del %s</div>"
        % run_time.strftime("%d/%m/%Y %H:%M:%S"),
        "<table class='kpi'><tr><td>File esaminati</td><td>Record esaminati</td>"
        "<td>Record errati</td><td>Record validi</td></tr>",
        "<tr><td class='v'>%d</td><td class='v'>%d</td><td class='v'>%d</td>"
        "<td class='v %s'>%s</td></tr></table>"
        % (len(results), tot, bad, pct_class(total_pct), fmt_pct(total_pct)),
        "<h2>Riepilogo file</h2>",
        "<table><thead><tr><th>File</th><th class='n'>Record</th>"
        "<th class='n'>Errati</th><th class='n'>Validi</th>"
        "<th class='n'>Segnalazioni</th><th>Modello Odoo</th></tr></thead><tbody>",
    ]
    for r in results:
        out.append(
            "<tr><td>%s</td><td class='n'>%d</td><td class='n'>%d</td>"
            "<td class='n %s'>%s</td><td class='n'>%d</td><td class='small'>%s</td></tr>"
            % (h(os.path.basename(r.path)), r.records, r.bad_records,
               pct_class(r.valid_pct), fmt_pct(r.valid_pct), len(r.issues),
               h(r.model or "-")))
    out.append("</tbody></table>")
    for title, header, rows in matches:
        out += ["<h2>%s</h2>" % h(title),
                "<div class='meta'>Record con la stessa partita IVA: %d</div>" % len(rows)]
        if rows:
            out.append("<table><thead><tr>%s</tr></thead><tbody>"
                       % "".join("<th>%s</th>" % h(c) for c in header))
            out += ["<tr>%s</tr>" % "".join("<td>%s</td>" % h(c) for c in row)
                    for row in rows]
            out.append("</tbody></table>")
    for r in results:
        out += [
            "<h2>%s</h2>" % h(os.path.basename(r.path)),
            "<div class='meta'><span class='small'>%s</span><br/>"
            "Record %d &middot; errati %d &middot; validi <span class='%s'>%s</span>"
            " &middot; <span class='small'>modello Odoo %s</span></div>"
            % (h(r.path), r.records, r.bad_records,
               pct_class(r.valid_pct), fmt_pct(r.valid_pct), h(r.model or "-")),
        ]
        if not r.issues:
            out.append("<div class='ok'>Nessuna segnalazione.</div>")
            continue
        out.append(
            "<table class='issues'><colgroup><col width='22%%'/><col width='14%%'/>"
            "<col width='64%%'/></colgroup>"
            # il nome del file è nell'intestazione della tabella, che viene
            # ripetuta in cima a ogni pagina ([subsection] di wkhtmltopdf non
            # è affidabile quando una sezione inizia a metà pagina)
            "<thead><tr><th colspan='3' class='file'>%s</th></tr>"
            "<tr><th>Riga (foglio di calcolo)</th><th>Campo</th>"
            "<th>Errore / valore</th></tr></thead><tbody>"
            % h(os.path.basename(r.path)))
        for field, message, count, voci, full, searched in group_issues(r.issues):
            if full is not None:  # un solo record: riga unica
                rows, _value, detail = voci[0]
                out.append(
                    "<tr class='one'><td class='rows'>%s</td><td>%s</td><td>%s%s</td></tr>"
                    % (pdf_rows(rows), h(field), h(full), pdf_hint(detail)))
                continue
            # sottotitolo: campo, messaggio e numero di record
            note = ("<div class='hint'>Ricerca alternativa per e-mail, partita IVA "
                    "e dominio: in corsivo i record trovati; senza nota nessuna "
                    "corrispondenza</div>" if searched else "")
            out.append(
                "<tr class='grp'><td colspan='3'><b>%s</b> &mdash; %s"
                "<span class='cnt'>%d record</span>%s</td></tr>"
                % (h(field), h(message), count, note))
            for rows, value, detail in voci:
                # elenchi lunghi su più righe di tabella: una cella che
                # attraversa il cambio pagina viene tagliata da wkhtmltopdf
                parts = rows.split(", ")
                for ix in range(0, len(parts), PDF_ROWS_PER_CELL):
                    chunk = pdf_rows(", ".join(parts[ix:ix + PDF_ROWS_PER_CELL]))
                    cell = ("<span class='small'>(segue)</span>" if ix else
                            "%s%s" % ("<span class='val'>%s</span>" % h(value) if value
                                      else "<span class='small'>(vuoto)</span>",
                                      pdf_hint(detail)))
                    out.append("<tr><td class='rows'>%s</td><td colspan='2'>%s</td></tr>"
                               % (chunk, cell))
        out.append("</tbody></table>")
    if skipped:
        out += ["<h2>File non esaminati</h2>",
                "<table><thead><tr><th>File</th><th>Motivo</th></tr></thead><tbody>"]
        out += ["<tr><td>%s</td><td>%s</td></tr>" % (h(f), h(m)) for f, m in skipped]
        out.append("</tbody></table>")
    out.append("</body></html>")

    with tempfile.NamedTemporaryFile(
            "w", suffix=".html", encoding="utf-8", delete=False) as fd:
        fd.write("\n".join(out))
        html = fd.name
    footer = ["--footer-font-size", "7", "--footer-right", "Pagina [page] di [topage]",
              "--footer-left", "Controllo CSV Odoo 18.0 - %s"
              % run_time.strftime("%d/%m/%Y %H:%M")]
    try:
        proc = subprocess.run(
            [wkhtmltopdf, "--quiet", "--encoding", "utf-8",
             "--page-size", "A4", "--orientation", "Landscape",
             "--margin-top", "12mm", "--margin-bottom", "14mm",
             "--margin-left", "10mm", "--margin-right", "10mm"]
            # indice con i numeri di pagina, generato dai titoli h1/h2
            + ["toc", "--toc-header-text", "Indice"] + footer
            + [html] + footer + [path],
            capture_output=True, text=True)
    finally:
        os.unlink(html)
    if proc.returncode != 0 and not os.path.isfile(path):
        raise SystemExit("wkhtmltopdf non riuscito: %s" % proc.stderr.strip())


WRITERS = {".md": write_md, ".xlsx": write_xlsx, ".ods": write_ods, ".pdf": write_pdf}


# ---------------------------------------------------------------------------
def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Controllo coerenza file CSV per importazione Odoo 18.0")
    parser.add_argument("-c", "--config",
                        help="file di configurazione (default: odoo_csv_check.ini "
                             "nella directory corrente o in quella dello script)")
    parser.add_argument("-o", "--output",
                        help="file report (.pdf, .md, .xlsx, .ods): sovrascrive la configurazione")
    parser.add_argument("-V", "--version", action="version", version=__version__)
    args = parser.parse_args(argv)

    config = args.config
    if not config:
        config = "odoo_csv_check.ini"
        if not os.path.isfile(config):
            config = os.path.join(
                os.path.dirname(os.path.abspath(__file__)), config)
    cfg = Config(config)
    run_time = datetime.now()
    report = (args.output or cfg.report_file) % {
        "datetime": run_time.strftime("%Y%m%d_%H%M%S")}
    writer = WRITERS.get(os.path.splitext(report)[1].lower())
    if not writer:
        raise SystemExit("Estensione report non supportata: %s (.pdf, .md, .xlsx, .ods)" % report)

    maps = Mappings(cfg)
    files, skipped, nonconforming = select_files(cfg)
    # i file usati solo come elenco di valori di [lookup_only] (assenti in
    # dirmap, es. Agenti.csv) non sono file da importare
    lookups = {}
    for name in lookup_sources(cfg, files):
        if not maps.resolve(name, cfg.prefix_match, cfg.alias)[0]:
            lookups[name] = files.pop(name)
            skipped.append((os.path.basename(lookups[name]),
                            "elenco di valori di [lookup_only]"))
    checker = Checker(cfg, maps, files, nonconforming)
    checker.lookup_files = lookups
    for name in sorted(files):  # [iban] prima delle chiavi: IBAN = counterpart_pk
        fixed, removed = checker.fix_iban(name, files[name])
        if fixed or removed:
            print("%-28s %d IBAN ricalcolati, %d righe rimosse da %s" % (
                "[iban]:", fixed, removed, os.path.basename(files[name])))
    checker.collect_keys()
    results = [checker.check_file(name, files[name]) for name in sorted(files)]
    for result in results:
        removed = checker.purge_file(result)
        if removed:
            print("%-28s %d record rimossi da %s" % (
                "[purge]:", removed, os.path.basename(result.path)))
    checker.add_insert_issues(results)
    matches = checker.vat_matches()
    writer(report, run_time, results, skipped, matches)

    for label, value in summary_rows(run_time, results):
        print("%-28s %s" % (label + ":", value))
    for title, _header, rows in matches:
        print("%-28s %d" % (title + ":", len(rows)))
    print("Report: %s" % report)
    return 1 if any(r.issues for r in results) else 0


if __name__ == "__main__":
    sys.exit(main())
